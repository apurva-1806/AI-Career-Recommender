"""Plain-text extraction from uploaded resumes (.pdf, .docx, .txt, .md)."""
from __future__ import annotations

import io
from pathlib import PurePath

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class ResumeError(ValueError):
    """Raised when a resume cannot be read."""


def extract_text(filename: str, payload: bytes) -> str:
    extension = PurePath(filename or "").suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeError(f"Unsupported file type '{extension or 'unknown'}'. "
                          f"Upload a PDF, DOCX, TXT or MD file.")
    if not payload:
        raise ResumeError("The uploaded file is empty.")

    try:
        if extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(payload))
            if reader.is_encrypted:
                raise ResumeError("This PDF is password-protected. Remove the password and upload it again.")
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif extension == ".docx":
            from docx import Document

            document = Document(io.BytesIO(payload))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        else:
            text = payload.decode("utf-8", errors="replace")
    except ResumeError:
        raise
    except Exception as exc:
        raise ResumeError(f"Could not read this {extension[1:].upper()} file. "
                          f"Try exporting it again or paste the text instead.") from exc

    text = "\n".join(line.strip() for line in text.replace("\x00", "").splitlines() if line.strip())
    if len(text) < 20:
        raise ResumeError("No readable text was found. Scanned PDFs need OCR first; paste the text instead.")
    return text
